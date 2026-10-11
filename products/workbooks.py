"""Editable XLSX workbooks for DeskNest Studio products (openpyxl)."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, INK, MUTED = "1F2A44", "2B2F36", "6B7280"
LINE, ALT, CALC, INPUT = "E3E6EB", "F8F9FB", "EEF0F4", "FFF8E1"
GOOD_FILL, GOOD_FONT = "E6F2EA", "1E6B43"
WARN_FILL, WARN_FONT = "FFF2DB", "8A5A00"
BAD_FONT = "B42318"
DATE_FMT = "dd mmm yyyy"
YEAR = dt.date.today().year
SITE = {}  # filled in by generate.py

THIN = Side(style="thin", color=LINE)
HAIR = Side(style="thin", color="D5D9E0")


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


class Col:
    """A tracker column.

    kind: text | date | int | num | pct | list | formula | static
    """

    def __init__(self, name, width=16, kind="text", options=None, formula=None, values=None, wrap=None,
                 fmt=None, strict=True):
        self.name, self.width, self.kind = name, width, kind
        self.options, self.formula, self.values = options, formula, values
        self.wrap = wrap if wrap is not None else width >= 24
        self.fmt, self.strict = fmt, strict


def sheet_title(ws, title, subtitle):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color=MUTED)
    ws["A2"].alignment = Alignment(vertical="top")


def input_cell(ws, row, label, kind="date", note=None):
    """A labelled input on the given row: label in A, value in C. Returns the absolute ref."""
    ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=INK)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=3)
    cell.fill = fill(INPUT)
    cell.border = Border(top=HAIR, bottom=HAIR, left=HAIR, right=HAIR)
    cell.font = Font(bold=True, color=NAVY)
    if kind == "date":
        cell.number_format = DATE_FMT
        dv = DataValidation(type="date", operator="greaterThan", formula1="1", errorStyle="warning",
                            error="Please enter a date, for example 2026-01-05.", errorTitle="Date expected")
        ws.add_data_validation(dv)
        dv.add(cell)
    if note:
        ws.cell(row=row, column=4, value=note).font = Font(italic=True, size=9, color=MUTED)
    return f"$C${row}"


def summary_row(ws, row, items):
    """Small KPI cells: label then formula, left to right."""
    col = 1
    for label, formula, fmt in items:
        lc = ws.cell(row=row, column=col, value=label)
        lc.font = Font(size=9, color=MUTED)
        lc.alignment = Alignment(horizontal="right")
        vc = ws.cell(row=row, column=col + 1, value=formula)
        vc.font = Font(bold=True, color=NAVY)
        vc.fill = fill(CALC)
        if fmt:
            vc.number_format = fmt
        col += 2


def tracker(ws, cols, n_rows, header_row=4, status_col=None, done=("Done",), waiting=(), overdue=(),
            row_height=None, freeze_col=1, landscape=True):
    """Lay out a styled table with validation and conditional formatting.

    overdue: list of (date_col_index, statuses_that_clear_it) pairs (0-based column index).
    """
    first, last = header_row + 1, header_row + n_rows
    ws.row_dimensions[header_row].height = 30
    for ci, col in enumerate(cols, start=1):
        letter = get_column_letter(ci)
        ws.column_dimensions[letter].width = col.width
        h = ws.cell(row=header_row, column=ci, value=col.name)
        h.font = Font(bold=True, color="FFFFFF", size=10)
        h.fill = fill(NAVY if col.kind != "formula" else "3B4A6B")
        h.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center" if col.width <= 12 else "left")
        h.border = Border(bottom=THIN)
        rng = f"{letter}{first}:{letter}{last}"
        for r in range(first, last + 1):
            c = ws.cell(row=r, column=ci)
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=col.wrap,
                                    horizontal="center" if col.kind in ("int", "static") and col.width <= 10 else None)
            if (r - first) % 2 == 1:
                c.fill = fill(ALT)
            if col.kind == "formula":
                c.value = col.formula.format(r=r)
                c.fill = fill(CALC)
                c.font = Font(color=MUTED)
            elif col.kind == "static" and col.values:
                c.value = col.values[r - first] if r - first < len(col.values) else None
                c.font = Font(color=MUTED)
            fmt = col.fmt or {"date": DATE_FMT, "pct": "0%", "num": "#,##0.##", "int": "0"}.get(col.kind)
            if fmt:
                c.number_format = fmt
        if col.kind == "list":
            dv = DataValidation(type="list", formula1='"' + ",".join(col.options) + '"', allow_blank=True,
                                showErrorMessage=col.strict, errorStyle="warning",
                                error="Choose a value from the list (or keep your own text).", errorTitle=col.name)
            ws.add_data_validation(dv)
            dv.add(rng)
        elif col.kind == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True,
                                errorStyle="warning", error="Please enter a date.", errorTitle="Date expected")
            ws.add_data_validation(dv)
            dv.add(rng)
        elif col.kind == "pct":
            dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                                errorStyle="warning", error="Enter a percentage between 0% and 100%.")
            ws.add_data_validation(dv)
            dv.add(rng)
            ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num",
                                                           end_value=1, color="8FA3C7"))
    last_letter = get_column_letter(len(cols))
    full = f"A{first}:{last_letter}{last}"
    if status_col is not None:
        s = get_column_letter(status_col + 1)
        cond = "OR(" + ",".join(f'${s}{first}="{v}"' for v in done) + ")"
        ws.conditional_formatting.add(f"{s}{first}:{s}{last}", FormulaRule(
            formula=[cond], fill=fill(GOOD_FILL), font=Font(color=GOOD_FONT, bold=True)))
        if waiting:
            cond = "OR(" + ",".join(f'${s}{first}="{v}"' for v in waiting) + ")"
            ws.conditional_formatting.add(f"{s}{first}:{s}{last}", FormulaRule(
                formula=[cond], fill=fill(WARN_FILL), font=Font(color=WARN_FONT)))
    for date_idx, clears in overdue:
        d = get_column_letter(date_idx + 1)
        clear = ""
        if status_col is not None and clears:
            s = get_column_letter(status_col + 1)
            clear = ",NOT(OR(" + ",".join(f'${s}{first}="{v}"' for v in clears) + "))"
        ws.conditional_formatting.add(f"{d}{first}:{d}{last}", FormulaRule(
            formula=[f'AND(${d}{first}<>"",${d}{first}<TODAY(){clear})'], font=Font(color=BAD_FONT, bold=True)))
    if row_height:
        for r in range(first, last + 1):
            ws.row_dimensions[r].height = row_height
    ws.freeze_panes = ws.cell(row=first, column=freeze_col)
    ws.auto_filter.ref = f"A{header_row}:{last_letter}{last}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = f"DeskNest Studio · {ws.title}"
    ws.oddFooter.right.text = "Page &P of &N"
    return first, last, full


def start_here(wb, title, description, sheets, tips):
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    ws["A1"] = "DeskNest Studio"
    ws["A1"].font = Font(bold=True, color=NAVY, size=11)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, color=NAVY, size=20)
    ws.row_dimensions[2].height = 32
    ws["A3"] = description
    ws["A3"].font = Font(color=INK, size=11)
    ws.merge_cells("A3:B3")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 34
    r = 5
    ws.cell(row=r, column=1, value="Sheet").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=2, value="What it's for").font = Font(bold=True, color="FFFFFF")
    for c in (1, 2):
        ws.cell(row=r, column=c).fill = fill(NAVY)
    for name, purpose in sheets:
        r += 1
        cell = ws.cell(row=r, column=1, value=name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(color=NAVY, underline="single", bold=True)
        pc = ws.cell(row=r, column=2, value=purpose)
        pc.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (1, 2):
            ws.cell(row=r, column=c).border = Border(bottom=THIN)
    r += 2
    ws.cell(row=r, column=1, value="Tips").font = Font(bold=True, color=NAVY, size=12)
    general = [
        "Yellow cells are inputs. Grey columns are calculated automatically, so there is no need to type in them.",
        "Drop-down lists are set up for status and category columns. You can type your own value if you prefer; "
        "you'll see a gentle warning, which you can accept.",
        "Dates turn red when they are in the past and the item is not finished yet.",
        "Use the filter arrows in each header row to sort or show only what you need.",
        "Works in Microsoft Excel, Google Sheets, LibreOffice Calc and Apple Numbers.",
    ]
    for t in tips + general:
        r += 1
        ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
        c = ws.cell(row=r, column=2, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
    ws.cell(row=r, column=1, value="License").font = Font(bold=True, color=NAVY, size=12)
    r += 1
    c = ws.cell(row=r, column=2, value=(
        "You may use and edit this workbook for your own personal or internal business purposes. Please do not "
        f"resell, share or redistribute it. Full terms: https://{SITE['domain']}/terms/"))
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 2
    ws.cell(row=r, column=1, value="Support").font = Font(bold=True, color=NAVY, size=12)
    r += 1
    ws.cell(row=r, column=2, value=(
        f"Questions or a problem with this file? Email {SITE['email']} with your order reference."))
    r += 2
    c = ws.cell(row=r, column=2, value=(
        "This workbook is an organizational tool. It is not legal, financial, accounting, tax, investment, "
        "medical or other professional advice."))
    c.font = Font(italic=True, size=9, color=MUTED)
    c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 26
    r += 1
    ws.cell(row=r, column=2, value=f"© {YEAR} DeskNest Studio").font = Font(size=9, color=MUTED)
    ws.sheet_properties.tabColor = NAVY


def new_sheet(wb, title, subtitle):
    ws = wb.create_sheet(title)
    sheet_title(ws, title, subtitle)
    ws.sheet_properties.tabColor = "8FA3C7"
    return ws


def finish(wb, path, title):
    wb.properties.creator = "DeskNest Studio"
    wb.properties.title = f"{title} — DeskNest Studio"
    wb.properties.lastModifiedBy = "DeskNest Studio"
    wb.active = 0
    wb.save(path)


WEEKDAY = 'CHOOSE(WEEKDAY({d},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")'
WEEKEND_CF = 'AND(${d}{r}<>"",WEEKDAY(${d}{r},2)>5)'


def _calendar_sheet(wb, title, subtitle, days, extra_cols, status_col_offset, done):
    ws = new_sheet(wb, title, subtitle)
    start = input_cell(ws, 3, "Start date", note="Enter the first day of your plan; all dates fill in automatically.")
    cols = [Col("Day", 6, "static", values=list(range(1, days + 1))),
            Col("Date", 13, "formula", formula=f'=IF({start}="","",{start}+A{{r}}-1)', fmt=DATE_FMT),
            Col("Weekday", 9, "formula", formula='=IF(B{r}="","",' + WEEKDAY.format(d="B{r}") + ")"),
            Col("Week", 7, "formula", formula="=INT((A{r}-1)/7)+1", fmt="0")] + extra_cols
    first, last, full = tracker(ws, cols, days, header_row=5, status_col=4 + status_col_offset, done=done,
                                freeze_col=4)
    ws.conditional_formatting.add(full, FormulaRule(formula=[WEEKEND_CF.format(d="B", r=first)], fill=fill("F1F2F5")))
    return ws


# --------------------------------------------------------------------------
# Workbooks
# --------------------------------------------------------------------------

FORMATS_GENERAL = ["Article", "Post", "Video", "Newsletter", "Podcast", "Image", "Other"]
CONTENT_STATUS = ["Idea", "Planned", "Drafting", "Review", "Scheduled", "Published"]
TASK_STATUS = ["Not started", "In progress", "Waiting", "Done"]
PRIORITY = ["High", "Medium", "Low"]


def content_planning_calendar(path):
    wb = Workbook()
    start_here(wb, "Content Planning Calendar",
               "A 90-day system for planning content ideas, publishing schedules and campaign tasks.",
               [("90-Day Calendar", "Enter a start date and plan one piece of content per day or as often as suits you."),
                ("Content Ideas", "Collect and prioritize ideas before they go into the calendar."),
                ("Publishing Tracker", "Track each piece from draft to published, with due dates and links."),
                ("Campaign Plan", "Plan one campaign: objective, audience, deliverables and tasks."),
                ("Content Review", "Review published content and decide what to keep, change or stop.")],
               ["Start on the 90-Day Calendar sheet: type your start date in the yellow cell."])
    _calendar_sheet(wb, "90-Day Calendar", "Plan your next 90 days. Weekends are shaded.", 90,
                    [Col("Channel", 16), Col("Content / working title", 40), Col("Format", 14, "list", FORMATS_GENERAL),
                     Col("Owner", 14), Col("Status", 13, "list", CONTENT_STATUS), Col("Notes", 30)],
                    status_col_offset=4, done=("Published",))

    ws = new_sheet(wb, "Content Ideas", "Capture ideas as they come. Sort by priority when you plan the next block.")
    tracker(ws, [Col("Idea / working title", 40), Col("Who it's for", 22), Col("Format", 14, "list", FORMATS_GENERAL),
                 Col("Channel", 16), Col("Priority", 11, "list", PRIORITY), Col("Date added", 13, "date"),
                 Col("In calendar?", 12, "list", ["Yes", "No"]), Col("Notes", 30)], 80, status_col=6, done=("Yes",))

    ws = new_sheet(wb, "Publishing Tracker", "Every piece from draft to published. Overdue dates turn red.")
    tracker(ws, [Col("Title", 36), Col("Channel", 16), Col("Format", 14, "list", FORMATS_GENERAL), Col("Owner", 14),
                 Col("Draft due", 13, "date"), Col("Publish date", 13, "date"),
                 Col("Status", 13, "list", CONTENT_STATUS),
                 Col("Days to publish", 11, "formula",
                     formula='=IF(OR(F{r}="",G{r}="Published"),"",F{r}-TODAY())', fmt="0"),
                 Col("Link / location", 28), Col("Notes", 28)], 100, status_col=6, done=("Published",),
            waiting=("Review", "Scheduled"),
            overdue=[(4, ("Review", "Scheduled", "Published")), (5, ("Published",))])

    ws = new_sheet(wb, "Campaign Plan", "One campaign per copy of this sheet. Right-click the tab to duplicate it.")
    labels = ["Campaign name", "Objective", "Audience", "Key message", "Call to action", "Start date", "End date",
              "Owner", "How results will be reviewed"]
    for i, lab in enumerate(labels):
        r = 4 + i
        ws.cell(row=r, column=1, value=lab).font = Font(bold=True, color=INK)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2)
        c.fill = fill(INPUT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if "date" in lab:
            c.number_format = DATE_FMT
        ws.row_dimensions[r].height = 22
    ws.cell(row=14, column=1, value="Deliverables").font = Font(bold=True, color=NAVY, size=12)
    tracker(ws, [Col("Deliverable", 34), Col("Channel", 16), Col("Owner", 14), Col("Due", 13, "date"),
                 Col("Status", 13, "list", CONTENT_STATUS), Col("Notes", 28)], 15, header_row=15, status_col=4,
            done=("Published",), overdue=[(3, ("Scheduled", "Published"))])
    ws.cell(row=32, column=1, value="Tasks").font = Font(bold=True, color=NAVY, size=12)
    tracker(ws, [Col("Task", 34), Col("Owner", 16), Col("Due", 14, "date"), Col("Done?", 13, "list", ["Yes", "No"])],
            15, header_row=33, status_col=3, done=("Yes",), overdue=[(2, ("Yes",))])
    ws.freeze_panes = None

    ws = new_sheet(wb, "Content Review", "Rename the Measure headers to the numbers you track (for example views or sign-ups).")
    tracker(ws, [Col("Content", 36), Col("Date published", 14, "date"), Col("Channel", 16), Col("Measure 1", 12, "num"),
                 Col("Measure 2", 12, "num"), Col("Measure 3", 12, "num"),
                 Col("Decision", 12, "list", ["Keep", "Change", "Stop"]), Col("What worked", 30),
                 Col("What to change", 30)], 60, status_col=6, done=("Keep",), waiting=("Change",))
    finish(wb, path, "Content Planning Calendar")


def productivity_toolkit_trackers(path):
    wb = Workbook()
    start_here(wb, "Productivity Toolkit Trackers",
               "Spreadsheet companions to the Complete Productivity Toolkit: priorities, project tasks, action items, "
               "decisions, weekly reviews and an ideal-week planner.",
               [("Priorities", "Rank what matters this week and track it to done."),
                ("Project Tasks", "Break projects into tasks with owners, dates, status and progress."),
                ("Action Items", "Collect action items from all your meetings in one list."),
                ("Decisions", "A running log of decisions, the reasons behind them and when to review them."),
                ("Weekly Review Log", "One line per weekly review, with simple 1–5 scores to spot patterns."),
                ("Ideal Week", "Block out a realistic week in 30-minute slots using categories.")],
               ["The summary cells at the top of some sheets update as you fill in the table."])

    ws = new_sheet(wb, "Priorities", "Rank your priorities and keep the next action small and specific.")
    tracker(ws, [Col("Rank", 7, "int"), Col("Priority", 36), Col("Why it matters", 30), Col("Next action", 30),
                 Col("Deadline", 13, "date"), Col("Status", 13, "list", TASK_STATUS),
                 Col("Completed on", 13, "date")], 40, status_col=5, waiting=("Waiting",),
            overdue=[(4, ("Done",))])

    ws = new_sheet(wb, "Project Tasks", "List every task with an owner and due date. Use the Project column to track several projects.")
    f, l = 6, 105
    summary_row(ws, 3, [("Tasks", f'=COUNTA(B{f}:B{l})', None), ("Done", f'=COUNTIF(G{f}:G{l},"Done")', None),
                        ("Overdue", f'=COUNTIFS(E{f}:E{l},"<"&TODAY(),G{f}:G{l},"<>Done")', None),
                        ("Complete", f'=IFERROR(COUNTIF(G{f}:G{l},"Done")/COUNTA(B{f}:B{l}),0)', "0%")])
    tracker(ws, [Col("Project", 18), Col("Task", 36), Col("Owner", 14), Col("Start", 13, "date"), Col("Due", 13, "date"),
                 Col("Priority", 11, "list", PRIORITY), Col("Status", 13, "list", TASK_STATUS),
                 Col("% complete", 11, "pct"),
                 Col("Days left", 10, "formula", formula='=IF(OR(E{r}="",G{r}="Done"),"",E{r}-TODAY())', fmt="0"),
                 Col("Notes", 28)], 100, header_row=5, status_col=6, waiting=("Waiting",), overdue=[(4, ("Done",))])

    ws = new_sheet(wb, "Action Items", "Every action item from every meeting. Review open items at the start of each meeting.")
    f, l = 6, 155
    summary_row(ws, 3, [("Open", f'=COUNTIFS(B{f}:B{l},"<>",G{f}:G{l},"<>Done")', None),
                        ("Overdue", f'=COUNTIFS(F{f}:F{l},"<"&TODAY(),G{f}:G{l},"<>Done")', None)])
    tracker(ws, [Col("No.", 6, "static", values=list(range(1, 151))), Col("Action item", 40), Col("Owner", 14),
                 Col("Meeting", 20), Col("Meeting date", 13, "date"), Col("Due date", 13, "date"),
                 Col("Status", 13, "list", ["Open", "In progress", "Blocked", "Done"]), Col("Notes", 30)], 150,
            header_row=5, status_col=6, waiting=("Blocked",), overdue=[(5, ("Done",))])

    ws = new_sheet(wb, "Decisions", "Record what was decided, why, and when it should be revisited.")
    tracker(ws, [Col("Date", 13, "date"), Col("Decision", 40), Col("Context & options considered", 40),
                 Col("Decided by", 16), Col("Who is affected", 22), Col("Review date", 13, "date")], 100)

    ws = new_sheet(wb, "Weekly Review Log", "One row per week. Scores are 1 (low) to 5 (high).")
    f, l = 6, 57
    summary_row(ws, 3, [("Avg focus", f'=IFERROR(AVERAGE(E{f}:E{l}),"")', "0.0"),
                        ("Avg energy", f'=IFERROR(AVERAGE(F{f}:F{l}),"")', "0.0"),
                        ("Avg progress", f'=IFERROR(AVERAGE(G{f}:G{l}),"")', "0.0")])
    first, last, _ = tracker(ws, [Col("Week of", 13, "date"), Col("Top 3 done?", 12, "list", ["Yes", "Partly", "No"]),
                                  Col("Biggest win", 30), Col("Main obstacle", 30), Col("Focus 1–5", 10, "int"),
                                  Col("Energy 1–5", 10, "int"), Col("Progress 1–5", 11, "int"),
                                  Col("Change for next week", 32)], 52, header_row=5, status_col=1, done=("Yes",),
                             waiting=("Partly",))
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        errorStyle="warning", error="Use a score from 1 to 5.")
    ws.add_data_validation(dv)
    dv.add(f"E{first}:G{last}")

    ws = new_sheet(wb, "Ideal Week", "Choose a category for each 30-minute slot. Colors update automatically.")
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    cats = {"Focus": ("DCE5F5", NAVY), "Meetings": ("FCE8D5", "7A3E00"), "Admin": ("E9E7F7", "3F347A"),
            "Personal": ("E2F1E7", GOOD_FONT), "Break": ("F2F3F5", MUTED)}
    times = [f"{h}:{m:02d}" for h in range(7, 22) for m in (0, 30)]
    tracker(ws, [Col("Time", 9, "static", values=times)] + [Col(d, 15, "list", list(cats)) for d in days],
            len(times), status_col=None, freeze_col=2)
    rng = f"B5:H{4 + len(times)}"
    for name, (bg, fg) in cats.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'B5="{name}"'], fill=fill(bg), font=Font(color=fg)))
    ws["J4"] = "Hours per category"
    ws["J4"].font = Font(bold=True, color=NAVY)
    for i, name in enumerate(cats):
        ws.cell(row=5 + i, column=10, value=name)
        ws.cell(row=5 + i, column=11, value=f'=COUNTIF({rng},"{name}")/2').number_format = "0.0"
    ws.column_dimensions["J"].width = 18
    finish(wb, path, "Productivity Toolkit Trackers")


def small_business_workflow_trackers(path):
    wb = Workbook()
    start_here(wb, "Small Business Workflow Trackers",
               "Spreadsheet companions to the Small Business Workflow Toolkit for documenting workflows, "
               "routine tasks, responsibilities and weekly operations.",
               [("Workflow Steps", "Write out each workflow step by step, with owners, tools and time needed."),
                ("Recurring Tasks", "Routine tasks with their frequency; the next due date is calculated for you."),
                ("Responsibilities", "Who does the work, who owns the result, who is consulted and who is informed."),
                ("Weekly Operations", "Plan the operational week day by day."),
                ("Task Review", "Review completed work, note issues and track improvements."),
                ("Change Log", "Record changes to your processes and documentation.")],
               ["On Recurring Tasks, enter the date a task was last done; the next due date and days remaining "
                "update automatically."])

    ws = new_sheet(wb, "Workflow Steps", "Use the Workflow column to keep several workflows in one sheet.")
    f, l = 6, 125
    summary_row(ws, 3, [("Steps", f"=COUNTA(C{f}:C{l})", None), ("Total minutes", f"=SUM(F{f}:F{l})", "0")])
    tracker(ws, [Col("Workflow", 20), Col("Step no.", 9, "int"), Col("Step", 38), Col("Who", 14),
                 Col("Tool / template", 20), Col("Time (min)", 11, "num"), Col("Notes", 30)], 120, header_row=5)

    ws = new_sheet(wb, "Recurring Tasks", "Days until due turns red when a task is overdue and amber when it is due within 3 days.")
    freq = ["Daily", "Weekly", "Every 2 weeks", "Monthly", "Quarterly", "Yearly"]
    match = "{{" + ",".join(f'"{x}"' for x in freq) + "}}"  # doubled braces survive str.format
    nxt = ('=IF(OR(E{r}="",B{r}=""),"",IFERROR(CHOOSE(MATCH(B{r},' + match +
           ',0),E{r}+1,E{r}+7,E{r}+14,EDATE(E{r},1),EDATE(E{r},3),EDATE(E{r},12)),""))')
    first, last, _ = tracker(ws, [Col("Task", 34), Col("Frequency", 14, "list", freq), Col("Owner", 14),
                                  Col("Backup", 14), Col("Last done", 13, "date"),
                                  Col("Next due", 13, "formula", formula=nxt, fmt=DATE_FMT),
                                  Col("Days until due", 11, "formula", formula='=IF(F{r}="","",F{r}-TODAY())', fmt="0"),
                                  Col("Notes", 30)], 80)
    g = f"G{first}:G{last}"
    ws.conditional_formatting.add(g, FormulaRule(formula=[f'AND(G{first}<>"",G{first}<0)'], font=Font(color=BAD_FONT, bold=True)))
    ws.conditional_formatting.add(g, FormulaRule(formula=[f'AND(G{first}<>"",G{first}>=0,G{first}<=3)'],
                                                 fill=fill(WARN_FILL), font=Font(color=WARN_FONT)))

    ws = new_sheet(wb, "Responsibilities", "Name one accountable person per row.")
    tracker(ws, [Col("Area or task", 34), Col("Responsible (does it)", 18), Col("Accountable (owns it)", 18),
                 Col("Consulted", 18), Col("Informed", 18), Col("Backup", 16), Col("Notes", 30)], 60)

    ws = new_sheet(wb, "Weekly Operations", "Copy this sheet for each week, or overwrite it every Monday.")
    start = input_cell(ws, 3, "Week starting (Monday)")
    tracker(ws, [Col("Day", 12, "static", values=["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]),
                 Col("Date", 13, "formula", formula=f'=IF({start}="","",{start}+ROW()-6)', fmt=DATE_FMT),
                 Col("Key tasks", 40), Col("Deadlines", 28), Col("Availability / who's out", 24), Col("Notes", 30)],
            5, header_row=5, row_height=72, freeze_col=1)
    ws["A12"] = "This week's priorities"
    ws["A12"].font = Font(bold=True, color=NAVY, size=12)
    tracker(ws, [Col("Priority", 52), Col("Owner", 16), Col("Done?", 13, "list", ["Yes", "No"])], 5, header_row=13,
            status_col=2, done=("Yes",))
    ws.freeze_panes = None

    ws = new_sheet(wb, "Task Review", "Agree on one improvement for each issue and track it to done.")
    tracker(ws, [Col("Date", 13, "date"), Col("Task or process", 30), Col("On time?", 10, "list", ["Yes", "No"]),
                 Col("Time taken (min)", 11, "num"), Col("What went well", 28), Col("Issue noticed", 28),
                 Col("Improvement", 28), Col("Owner", 14), Col("Improvement status", 13, "list", ["Open", "Done"])], 80,
            status_col=8, waiting=("Open",))

    ws = new_sheet(wb, "Change Log", "Record every change to a documented process.")
    tracker(ws, [Col("Process", 26), Col("Version", 10), Col("Date", 13, "date"), Col("What changed", 44),
                 Col("Changed by", 16), Col("Approved by", 16)], 80)
    finish(wb, path, "Small Business Workflow Trackers")


def creator_planning_trackers(path):
    wb = Workbook()
    creator_formats = ["Video", "Short video", "Post", "Carousel", "Article", "Newsletter", "Podcast", "Livestream", "Other"]
    creator_status = ["Idea", "Scripting", "Producing", "Editing", "Scheduled", "Published"]
    start_here(wb, "Creator Planning Trackers",
               "Spreadsheet companions to the Creator Planning Bundle: calendar, idea bank, production checklist, "
               "publishing tracker, campaigns and performance notes.",
               [("Content Calendar", "Enter a start date and plan the next three months day by day."),
                ("Idea Bank", "Every idea with its hook, format, effort and priority."),
                ("Production Checklist", "Tick off each production step per piece; progress is calculated."),
                ("Publishing Tracker", "Track planned and published dates and links across platforms."),
                ("Campaigns", "Overview of launches, series and promotions."),
                ("Campaign Content", "The content pieces for each campaign, by phase."),
                ("Performance Log", "Note the measures you track and what you learn from each piece.")],
               ["On Production Checklist, choose ✓ in each step column (or type any mark) to update progress."])

    _calendar_sheet(wb, "Content Calendar", "Plan three months of content. Weekends are shaded.", 92,
                    [Col("Platform", 16), Col("Working title", 38), Col("Format", 14, "list", creator_formats),
                     Col("Status", 13, "list", creator_status), Col("Notes", 30)],
                    status_col_offset=3, done=("Published",))

    ws = new_sheet(wb, "Idea Bank", "Capture every idea with a hook so you always have something ready to make.")
    tracker(ws, [Col("Idea", 36), Col("Hook / angle", 32), Col("Format", 14, "list", creator_formats),
                 Col("Platform", 16), Col("Effort", 10, "list", ["Small", "Medium", "Large"]),
                 Col("Priority", 10, "list", PRIORITY), Col("Date added", 13, "date"),
                 Col("Status", 11, "list", ["New", "Planned", "Used", "Parked"])], 100, status_col=7, done=("Used",),
            waiting=("Planned",))

    ws = new_sheet(wb, "Production Checklist", "One row per piece of content. Choose ✓ as each step is completed.")
    steps = ["Outline / script", "Assets ready", "Recorded / written", "Edited", "Captions", "Cover / thumbnail",
             "Alt text", "Proofread", "Scheduled", "Published"]
    first, last, _ = tracker(ws, [Col("Working title", 32), Col("Platform", 14), Col("Publish date", 13, "date")] +
                             [Col(s, 10, "list", ["✓"], strict=False) for s in steps] +
                             [Col("Progress", 11, "formula", formula="=COUNTA(D{r}:M{r})/10", fmt="0%")], 60)
    ws.conditional_formatting.add(f"D{first}:M{last}", FormulaRule(formula=[f'D{first}<>""'], fill=fill(GOOD_FILL),
                                                                    font=Font(color=GOOD_FONT, bold=True)))
    ws.conditional_formatting.add(f"N{first}:N{last}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=1, color="8FA3C7"))
    for ci in range(4, 14):
        for r in range(first, last + 1):
            ws.cell(row=r, column=ci).alignment = Alignment(horizontal="center", vertical="top")

    ws = new_sheet(wb, "Publishing Tracker", "Planned dates turn red when they pass before the piece is published.")
    tracker(ws, [Col("Title", 36), Col("Platform", 16), Col("Format", 14, "list", creator_formats),
                 Col("Status", 13, "list", creator_status), Col("Planned date", 13, "date"),
                 Col("Published date", 13, "date"), Col("Link", 30), Col("Notes", 28)], 120, status_col=3,
            done=("Published",), waiting=("Scheduled",), overdue=[(4, ("Published",))])

    ws = new_sheet(wb, "Campaigns", "One row per launch, series or promotion.")
    tracker(ws, [Col("Campaign", 26), Col("Goal", 32), Col("Audience", 22), Col("Start", 13, "date"),
                 Col("Main date", 13, "date"), Col("End", 13, "date"),
                 Col("Status", 12, "list", ["Planning", "Active", "Complete"]), Col("Notes", 30)], 30, status_col=6,
            done=("Complete",), waiting=("Active",))

    ws = new_sheet(wb, "Campaign Content", "List the pieces for each campaign by phase.")
    tracker(ws, [Col("Campaign", 24), Col("Phase", 14, "list", ["Before launch", "Launch week", "After launch"]),
                 Col("Content piece", 34), Col("Platform", 16), Col("Due", 13, "date"),
                 Col("Status", 13, "list", creator_status), Col("Owner", 14)], 100, status_col=5, done=("Published",),
            overdue=[(4, ("Scheduled", "Published"))])

    ws = new_sheet(wb, "Performance Log", "Rename the Measure headers to the numbers you track (for example views, saves or sign-ups).")
    tracker(ws, [Col("Content", 34), Col("Platform", 16), Col("Date published", 14, "date"), Col("Measure 1", 12, "num"),
                 Col("Measure 2", 12, "num"), Col("Measure 3", 12, "num"), Col("What worked", 32),
                 Col("Try next time", 32)], 100)
    finish(wb, path, "Creator Planning Trackers")

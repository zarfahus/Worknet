"""Render template documents to PDF (print-ready) and DOCX (editable)."""
import datetime as dt

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

import blocks as B
from blocks import (FAINT, FILL, INK, MUTED, NAVY, RULE, TINT, WHITE, InfoBox, Note, Page, docx_new_page, layout,
                    spaced, sw, wrap)

M = 34  # page margin (pt)
YEAR = dt.date.today().year
SITE = {}  # filled in by generate.py from src/data/site.json


def license_line():
    return f"© {YEAR} WorkNest Studio · Licensed for personal and internal business use · Not for resale"


# --------------------------------------------------------------------------
# Logo
# --------------------------------------------------------------------------

def draw_logo(c, x, y, size):
    """Logo mark: navy tile with a roof line over two text lines (a 'nest' for work).
    (x, y) is the bottom-left corner; geometry mirrors assets/logo.svg (24-unit grid)."""
    u = size / 24.0
    c.setFillColor(NAVY)
    c.roundRect(x, y, size, size, 5.5 * u, stroke=0, fill=1)
    c.setStrokeColor(WHITE)
    c.setLineWidth(2 * u)
    c.setLineCap(1)
    c.setLineJoin(1)
    p = c.beginPath()
    p.moveTo(x + 6 * u, y + size - 11.5 * u)
    p.lineTo(x + 12 * u, y + size - 6.5 * u)
    p.lineTo(x + 18 * u, y + size - 11.5 * u)
    c.drawPath(p, stroke=1, fill=0)
    c.line(x + 8 * u, y + size - 15 * u, x + 16 * u, y + size - 15 * u)
    c.line(x + 8 * u, y + size - 18.5 * u, x + 13.5 * u, y + size - 18.5 * u)
    c.setLineCap(0)


def logo_png(path, px=256):
    img = Image.new("RGBA", (px, px), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    u = px / 24.0
    d.rounded_rectangle([0, 0, px - 1, px - 1], radius=int(5.5 * u), fill=(0x1F, 0x2A, 0x44, 255))
    wdt = int(2 * u)
    d.line([(6 * u, 11.5 * u), (12 * u, 6.5 * u), (18 * u, 11.5 * u)], fill="white", width=wdt, joint="curve")
    for (x1, y1), (x2, y2) in (((8 * u, 15 * u), (16 * u, 15 * u)), ((8 * u, 18.5 * u), (13.5 * u, 18.5 * u))):
        d.line([(x1, y1), (x2, y2)], fill="white", width=wdt)
    for (cx, cy) in ((6 * u, 11.5 * u), (18 * u, 11.5 * u), (8 * u, 15 * u), (16 * u, 15 * u), (8 * u, 18.5 * u),
                     (13.5 * u, 18.5 * u)):
        d.ellipse([cx - wdt / 2, cy - wdt / 2, cx + wdt / 2, cy + wdt / 2], fill="white")
    img.save(path)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _frame(c, W, H, doc_title, page_no):
    draw_logo(c, M, H - M - 13, 13)
    c.setFont(B.SEMI, 8.5)
    c.setFillColor(NAVY)
    c.drawString(M + 19, H - M - 9.5, "WorkNest Studio")
    c.setFont(B.REG, 7.5)
    c.setFillColor(MUTED)
    c.drawRightString(W - M, H - M - 9.5, doc_title)
    B.hline(c, M, W - M, H - M - 21, FAINT, 0.8)
    c.setFont(B.REG, 6.8)
    c.setFillColor(MUTED)
    c.drawString(M, M - 10, license_line())
    c.drawRightString(W - M, M - 10, f"Page {page_no}")


def draw_page(c, page, doc_title, page_no):
    W, H = landscape(A4) if page.landscape else A4
    c.setPageSize((W, H))
    _frame(c, W, H, doc_title, page_no)
    w = W - 2 * M
    c.setFont(B.SEMI, 19)
    c.setFillColor(NAVY)
    c.drawString(M, H - M - 48, page.title)
    y = H - M - 54
    if page.intro:
        c.setFont(B.REG, 8.5)
        c.setFillColor(MUTED)
        for ln in wrap(page.intro, B.REG, 8.5, w):
            y -= 12
            c.drawString(M, y, ln)
    top = y - 14
    bottom = M + 6
    try:
        heights = layout(page.blocks, w, top - bottom)
    except ValueError as exc:
        raise ValueError(f"{doc_title} / {page.title}: {exc}") from None
    for blk, h in zip(page.blocks, heights):
        blk.draw(c, M, top, w, h)
        top -= h + B.GAP


def _cover_illustration(c, x, y, w, h):
    """Three overlapping template sheets: a checklist, a table and a calendar grid."""
    sheets = [(x + w * 0.10, y + h * 0.16, -7), (x + w * 0.36, y + h * 0.10, 0), (x + w * 0.60, y + h * 0.18, 6)]
    sw_, sh_ = w * 0.30, h * 0.78
    for i, (sx, sy, rot) in enumerate(sheets):
        c.saveState()
        c.translate(sx + sw_ / 2, sy + sh_ / 2)
        c.rotate(rot)
        c.setFillColor(B.HexColor("#D9DEE7"))
        c.roundRect(-sw_ / 2 + 3, -sh_ / 2 - 4, sw_, sh_, 5, stroke=0, fill=1)
        c.setFillColor(WHITE)
        c.setStrokeColor(FAINT)
        c.roundRect(-sw_ / 2, -sh_ / 2, sw_, sh_, 5, stroke=1, fill=1)
        lx, rx, ty = -sw_ / 2 + 12, sw_ / 2 - 12, sh_ / 2 - 16
        c.setFillColor(NAVY)
        c.roundRect(lx, ty - 4, sw_ * 0.45, 5, 2, stroke=0, fill=1)
        if i == 0:
            for k in range(7):
                yy = ty - 22 - k * 17
                B.checkbox(c, lx, yy - 3, 7)
                B.hline(c, lx + 12, rx - (k % 3) * 14, yy, RULE, 0.8)
            c.setStrokeColor(NAVY)
            c.setLineWidth(1.2)
            c.line(lx + 1.5, ty - 22, lx + 3.5, ty - 24.5)
            c.line(lx + 3.5, ty - 24.5, lx + 7, ty - 18.5)
        elif i == 1:
            c.setFillColor(TINT)
            c.rect(lx, ty - 26, rx - lx, 12, stroke=0, fill=1)
            for k in range(8):
                B.hline(c, lx, rx, ty - 26 - (k + 1) * 14, RULE, 0.6)
            for fx in (0.35, 0.7):
                B.hline(c, lx + (rx - lx) * fx, lx + (rx - lx) * fx, ty - 14, FAINT)
                c.setStrokeColor(FAINT)
                c.line(lx + (rx - lx) * fx, ty - 14, lx + (rx - lx) * fx, ty - 26 - 8 * 14)
        else:
            cw, chh = (rx - lx) / 4, 22
            for r in range(5):
                for k in range(4):
                    c.setStrokeColor(FAINT)
                    c.setFillColor(FILL if (r + k) % 5 == 0 else WHITE)
                    c.rect(lx + k * cw, ty - 18 - (r + 1) * chh, cw, chh, stroke=1, fill=1)
        c.restoreState()


def draw_cover(c, doc, contents):
    W, H = A4
    c.setPageSize((W, H))
    band_h = H * 0.44
    c.setFillColor(B.HexColor("#EEF1F6"))
    c.rect(0, H - band_h, W, band_h, stroke=0, fill=1)
    draw_logo(c, M + 6, H - M - 26, 22)
    c.setFont(B.SEMI, 12)
    c.setFillColor(NAVY)
    c.drawString(M + 36, H - M - 18, "WorkNest Studio")
    _cover_illustration(c, M + 30, H - band_h + 30, W - 2 * M - 60, band_h - 110)

    x = M + 6
    y = H - band_h - 54
    c.setFont(B.SEMI, 28)
    c.setFillColor(NAVY)
    for ln in wrap(doc.title, B.SEMI, 28, W - 2 * x):
        c.drawString(x, y, ln)
        y -= 34
    c.setFont(B.REG, 11)
    c.setFillColor(INK)
    y -= 2
    for ln in wrap(doc.subtitle, B.REG, 11, W - 2 * x - 40):
        c.drawString(x, y, ln)
        y -= 16
    y -= 8
    spaced(c, x, y, "  ·  ".join(doc.format_labels), B.MED, 8.5, MUTED, 0.4)

    y -= 34
    spaced(c, x, y, "CONTENTS", B.SEMI, 8, NAVY, 0.8)
    y -= 8
    for name, page_no in contents:
        y -= 17
        c.setFont(B.REG, 10)
        c.setFillColor(INK)
        c.drawString(x, y, name)
        c.drawRightString(W - x, y, str(page_no))
        c.setDash(1, 2.5)
        B.hline(c, x + sw(name, B.REG, 10) + 6, W - x - sw(str(page_no), B.REG, 10) - 6, y + 1, RULE, 0.6)
        c.setDash()

    c.setFont(B.REG, 7.5)
    c.setFillColor(MUTED)
    c.drawString(x, M + 12, f"Version {doc.version} · {dt.date.today():%B %Y} · {SITE['domain']} · {SITE['email']}")
    c.drawString(x, M, license_line())


def how_to_page(doc):
    tips = [InfoBox("Getting started", doc.howto, cols=1)]
    tips.append(InfoBox("Printing", [
        "Pages are A4. To print on US Letter paper, choose “Fit to page” or “Scale to fit” in "
        "your print settings.",
        "Print single pages as often as you need them, for example one daily page per day or one notes page "
        "per meeting.",
        "You can also fill in the PDF on a tablet or computer with any PDF app that supports annotations."]))
    editable = []
    if "DOCX" in doc.all_formats:
        editable.append("DOCX files open in Microsoft Word, Google Docs, LibreOffice Writer and Apple Pages. "
                        "Change headings, add rows or delete pages to suit the way you work.")
    if "XLSX" in doc.all_formats:
        editable.append("XLSX workbooks open in Microsoft Excel, Google Sheets, LibreOffice Calc and Apple "
                        "Numbers. Drop-down lists, dates and formulas are set up for you; grey columns are "
                        "calculated automatically.")
    if editable:
        tips.append(InfoBox("Editable files", editable))
    tips.append(InfoBox("Your licence", [
        "You may use these templates for your own personal or internal business purposes, and edit them for "
        "your own use.",
        "Please do not resell, share or redistribute the files, or publish them as your own templates.",
        f"Full terms: {SITE['domain']}/terms/"]))
    tips.append(InfoBox("Need help?", [
        f"If a file will not open, is incomplete or does not match its description, email {SITE['email']} "
        "with your order reference and we will help."]))
    tips.append(Note("These templates are organisational tools. They are not legal, financial, accounting, tax, "
                     "investment, medical or other professional advice."))
    return Page("How to use these templates", None, tips, section="How to use these templates")


def build_pdf(doc, path):
    pages = [how_to_page(doc)] + doc.pages
    contents, seen = [], set()
    for i, p in enumerate(pages):
        if p.section not in seen:
            seen.add(p.section)
            contents.append((p.section, i + 2))
    c = canvas.Canvas(path, pagesize=A4, pageCompression=1)
    c.setTitle(f"{doc.title} — WorkNest Studio")
    c.setAuthor("WorkNest Studio")
    c.setSubject(doc.subtitle)
    c.setCreator("WorkNest Studio")
    c.setKeywords(["template", "planner", "WorkNest Studio"])
    draw_cover(c, doc, contents)
    c.showPage()
    for i, p in enumerate(pages):
        draw_page(c, p, doc.title, i + 2)
        c.showPage()
    c.save()
    return len(pages) + 1


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _page_number_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def _setup_docx(doc_title):
    d = DocxDocument()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.4)
        s.header_distance = s.footer_distance = Cm(0.7)
    normal = d.styles["Normal"]
    normal.font.name = B.DOCX_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), B.DOCX_FONT)
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(B.INK_HEX)
    normal.paragraph_format.space_after = Pt(2)
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 12)):
        st = d.styles[name]
        st.font.name = B.DOCX_FONT
        # Heading styles use theme fonts by default; pin them to the document font.
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:ascii"), B.DOCX_FONT)
        rfonts.set(qn("w:hAnsi"), B.DOCX_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(B.NAVY_HEX)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(4)
    # Title style has a bottom border in the default template; remove it.
    pPr = d.styles["Title"].element.pPr
    if pPr is not None and pPr.find(qn("w:pBdr")) is not None:
        pPr.remove(pPr.find(qn("w:pBdr")))

    hp = sec.header.paragraphs[0]
    r = hp.add_run("WorkNest Studio")
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(B.NAVY_HEX)
    r = hp.add_run(f"   |   {doc_title}")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(B.MUTED_HEX)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(license_line() + "   ·   Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(B.MUTED_HEX)
    _page_number_field(fp)
    return d


def build_docx(doc, path, logo_path):
    d = _setup_docx(doc.title)
    d.core_properties.title = f"{doc.title} — WorkNest Studio"
    d.core_properties.author = "WorkNest Studio"
    d.core_properties.subject = doc.subtitle
    d.core_properties.comments = license_line()

    d.add_picture(logo_path, width=Cm(1.3))
    p = d.add_paragraph(doc.title, style="Title")
    p.paragraph_format.space_before = Pt(10)
    p = d.add_paragraph()
    r = p.add_run(doc.subtitle)
    r.font.size = Pt(11.5)
    p = d.add_paragraph()
    r = p.add_run("Editable version  ·  " + "  ·  ".join(doc.format_labels))
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(B.MUTED_HEX)
    B.label_paragraph(d, "Contents")
    seen = []
    for pg in doc.pages:
        if pg.section not in seen:
            seen.append(pg.section)
    for name in seen:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(name).font.size = Pt(10.5)
    how = how_to_page(doc)
    p = d.add_paragraph(how.title, style="Heading 2")
    p.paragraph_format.space_before = Pt(14)
    for blk in how.blocks:
        blk.docx(B.body_ctx(d))
        B.small_gap(d, 6)

    for pg in doc.pages:
        docx_new_page(d, pg.landscape)
        d.add_paragraph(pg.title, style="Heading 1")
        if pg.intro:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(pg.intro)
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor.from_string(B.MUTED_HEX)
        ctx = B.body_ctx(d)
        for i, blk in enumerate(pg.blocks):
            if i:
                B.small_gap(ctx, 6)  # keeps adjacent tables from merging into one
            blk.docx(ctx)
    d.save(path)

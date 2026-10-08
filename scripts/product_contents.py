#!/usr/bin/env python3
"""Write src/data/contents.json: the page-by-page contents of every product PDF.

Computed from the same page definitions products/generate.py renders, so the product pages
list exactly what is inside each file. Does not touch the product files themselves.
"""
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "products"))
import generate  # noqa: E402
from render import how_to_page  # noqa: E402


def doc_contents(doc):
    pages = [how_to_page(doc)] + doc.pages
    sections = []
    for i, p in enumerate(pages):
        n = i + 2  # page 1 is the cover
        if sections and sections[-1]["name"] == p.section:
            sections[-1]["last"] = n
        else:
            sections.append({"name": p.section, "first": n, "last": n})
    for s in sections:
        s["pages"] = str(s["first"]) if s["first"] == s["last"] else f"{s['first']}–{s['last']}"
        del s["first"], s["last"]
    return {"title": doc.title, "pages": len(pages) + 1,
            "orientation": sorted({"landscape" if p.landscape else "portrait" for p in doc.pages}),
            "sections": [{"name": "Cover and contents", "pages": "1"}] + sections}


def workbook_sheets(filename):
    """Sheet names of a generated workbook (read from products/out, which generate.py writes)."""
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(generate.FILES_DIR, filename), read_only=True)
    return {"file": filename, "sheets": wb.sheetnames}


def main():
    docs = {d.key: d for d in generate.DOCS}
    out = {}
    for slug, file_ids in generate.PRODUCT_FILES.items():
        keys = []
        for fid in file_ids:
            key = fid.rsplit("-", 1)[0]
            if key in docs and fid.endswith("-pdf") and key not in keys:
                keys.append(key)
        out[slug] = {
            "pdfs": [dict(doc_contents(docs[k]), file=f"{docs[k].base}.pdf") for k in keys],
            "workbooks": [workbook_sheets(fn) for fid, fn, _, _ in generate.WORKBOOKS if fid in file_ids],
        }
    path = os.path.join(ROOT, "src", "data", "contents.json")
    with open(path, "w") as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    print(f"wrote {path}")


if __name__ == "__main__":
    main()

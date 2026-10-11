"""Editable XLSX workbooks for the Project Workflow Bundle.

Design rules, so the files behave the same in Excel, Google Sheets and LibreOffice:
- standard functions only (IF, COUNTIF(S), INDEX, MATCH, WEEKDAY, WEEKNUM, EOMONTH, DATE, TODAY)
- no macros, no TEXT() with locale-specific format codes (number formats are used instead)
- white cells are for typing; grey cells calculate; yellow cells are settings
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

GREEN = "2E5A44"
GREEN_TINT = "E3ECE4"
CALC_FILL = PatternFill("solid", fgColor="EFF1EF")
SETTING_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
TINT_FILL = PatternFill("solid", fgColor=GREEN_TINT)
THIN = Side(style="thin", color="D5DAD6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Arial"
DATE = "mm-dd-yy"  # built-in short date: shown in the user's regional date format

STATUS = ["Not started", "In progress", "Waiting", "On hold", "Done"]
PRIORITY = ["High", "Medium", "Low"]


def f(size=10, bold=False, color="1E2622", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def title(ws, text, hint, width_cols=10):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=GREEN)
    ws["A2"] = hint
    ws["A2"].font = f(9.5, color="4A554F", italic=True)
    ws["A2"].alignment = Alignment(wrap_text=False, vertical="top")
    ws.row_dimensions[1].height = 24


def header(ws, row, names, widths):
    for i, (name, width) in enumerate(zip(names, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = f(10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 22


def body(ws, first, last, ncols, calc_cols=(), date_cols=(), formats=None):
    formats = formats or {}
    for r in range(first, last + 1):
        for col in range(1, ncols + 1):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = f(10)
            c.alignment = Alignment(vertical="top", wrap_text=col not in date_cols)
            if col in calc_cols:
                c.fill = CALC_FILL
                c.font = f(10, color="4A554F")
            if col in date_cols:
                c.number_format = DATE
            if col in formats:
                c.number_format = formats[col]


def validation(ws, formula, ranges, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.error = "Please choose an option from the list. You can change the list on the Lists sheet."
    dv.errorTitle = "Not in the list"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def fill_rule(ws, rng, formula, fill=None, font=None):
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=fill, bgColor=fill) if fill else None, font=font),
    )


def print_setup(ws, area, title_rows=None, landscape=True):
    ws.print_area = area
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    if title_rows:
        ws.print_title_rows = title_rows
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def start_sheet(ws, heading, intro, steps, legend, extra):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    ws["B2"] = heading
    ws["B2"].font = Font(name=FONT, size=18, bold=True, color=GREEN)
    ws["B3"] = "ClearWeek Studio · evadeclaw.com · support@evadeclaw.com"
    ws["B3"].font = f(9.5, color="4A554F")
    r = 5
    ws.cell(row=r, column=2, value=intro).font = f(11)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 46
    r += 2
    ws.cell(row=r, column=2, value="Getting started").font = f(12, bold=True, color=GREEN)
    for i, step in enumerate(steps, start=1):
        r += 1
        ws.cell(row=r, column=1, value=i).font = f(11, bold=True, color=GREEN)
        c = ws.cell(row=r, column=2, value=step)
        c.font = f(11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32 if len(step) > 95 else 18
    r += 2
    ws.cell(row=r, column=2, value="Cell colors").font = f(12, bold=True, color=GREEN)
    for fill, text in legend:
        r += 1
        ws.cell(row=r, column=1).fill = fill
        ws.cell(row=r, column=1).border = BOX
        ws.cell(row=r, column=2, value=text).font = f(11)
    r += 2
    for heading_text, text in extra:
        ws.cell(row=r, column=2, value=heading_text).font = f(12, bold=True, color=GREEN)
        r += 1
        c = ws.cell(row=r, column=2, value=text)
        c.font = f(11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 46
        r += 2
    print_setup(ws, f"A1:B{r}", landscape=False)


LEGEND = [
    (PatternFill(fill_type=None), "White: type here."),
    (CALC_FILL, "Grey: calculated automatically. Please don't type over these cells."),
    (SETTING_FILL, "Yellow: a setting you can change, such as the month to show."),
]

COMPAT = (
    "Works with",
    "Microsoft Excel 2016 or later and Microsoft 365, Excel for the web, Google Sheets (use File › Save as Google Sheets) and LibreOffice Calc 7 or later. "
    "Apple Numbers opens the file and keeps the formulas, but doesn't support the color rules, so highlights and timeline bars won't show. No macros are used.",
)
LICENSE = (
    "License",
    "For your personal use and internal work within your own business or organization. Please don't share the file outside your organization, resell it or publish it. "
    "Full terms: evadeclaw.com/terms. Help: support@evadeclaw.com.",
)


def lists_sheet(ws, columns):
    """columns: list of (heading, values, slots, note)"""
    title(ws, "Lists", "The options in the drop-down lists. Edit or add entries in the white cells.")
    ws.sheet_view.showGridLines = False
    refs = {}
    for i, (name, values, slots, note) in enumerate(columns):
        col = 1 + i * 2
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 22
        ws.column_dimensions[get_column_letter(col + 1)].width = 3
        c = ws.cell(row=4, column=col, value=name)
        c.font = f(10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.border = BOX
        for r in range(5, 5 + slots):
            cell = ws.cell(row=r, column=col, value=values[r - 5] if r - 5 < len(values) else None)
            cell.border = BOX
            cell.font = f(10)
        if note:
            n = ws.cell(row=5 + slots + 1, column=col, value=note)
            n.font = f(8.5, color="4A554F", italic=True)
            n.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[5 + slots + 1].height = 48
        refs[name] = f"=Lists!${letter}$5:${letter}${4 + slots}"
    return refs


def month_grid(ws, setting_cell, top, cols=7, day_rows=None):
    """Write a 6-week calendar grid starting at row `top` (header) and return the date-cell coordinates."""
    names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    for i, n in enumerate(names, start=1):
        c = ws.cell(row=top, column=i, value=n)
        c.font = f(10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = 19
    first = f"DATE(YEAR({setting_cell}),MONTH({setting_cell}),1)"
    cells = []
    for w in range(6):
        r = top + 1 + w * 3
        for d in range(cols):
            col = d + 1
            c = ws.cell(row=r, column=col, value=f"={first}-WEEKDAY({first},3)+{7 * w + d}")
            c.number_format = "d"
            c.font = f(11, bold=True, color=GREEN)
            c.fill = CALC_FILL
            c.alignment = Alignment(horizontal="left", vertical="top")
            c.border = Border(left=THIN, right=THIN, top=THIN)
            for k in (1, 2):
                cc = ws.cell(row=r + k, column=col)
                cc.fill = CALC_FILL
                cc.font = f(9, color="1E2622")
                cc.alignment = Alignment(wrap_text=True, vertical="top")
                cc.border = Border(left=THIN, right=THIN, bottom=THIN if k == 2 else None)
            cells.append((r, col))
        ws.row_dimensions[r].height = 17
        ws.row_dimensions[r + 1].height = 30
        ws.row_dimensions[r + 2].height = 28
        rng = f"A{r}:{get_column_letter(cols)}{r}"
        fill_rule(ws, rng, f"A{r}=TODAY()", fill="FFF2CC")
        fill_rule(ws, rng, f"MONTH(A{r})<>MONTH({setting_cell})", font=Font(color="B4BBB6"))
    return first, cells


# ---------------------------------------------------------------------------------------


P_FIRST, P_LAST = 5, 34
T_FIRST, T_LAST = 5, 304


def project_tracker(path: Path, sample: bool = False) -> list[str]:
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    start = wb.active
    start.title = "Start here"
    projects = wb.create_sheet("Projects")
    tasks = wb.create_sheet("Tasks")
    timeline = wb.create_sheet("Timeline")
    month = wb.create_sheet("Monthly overview")
    weekly = wb.create_sheet("Weekly checklist")
    lists = wb.create_sheet("Lists")

    refs = lists_sheet(
        lists,
        [
            ("Status", STATUS, 10, "Keep one status called Done: the tracker counts it as finished."),
            ("Priority", PRIORITY, 6, None),
            ("People", ["Me"], 20, "Add the names of everyone who owns projects or tasks."),
        ],
    )

    start_sheet(
        start,
        "Project Tracker",
        "Keep your projects and their tasks in one place, see them on a timeline and a monthly calendar, and run a short weekly review. "
        "Everything on the other sheets updates from what you type on the Projects and Tasks sheets.",
        [
            "Go to the Lists sheet and add the names of the people who own work (or keep “Me”).",
            "On the Projects sheet, add each project with an owner, status, priority, start date and due date.",
            "On the Tasks sheet, add the tasks for each project. Choose the project from the drop-down list so the task is counted.",
            "When a task is finished, set its status to Done and type the date in “Done on”. Progress and counts update automatically.",
            "Use the Timeline sheet to see projects week by week, and the Monthly overview to see what's due in any month.",
            "Once a week, work through the Weekly checklist and tick each step.",
        ],
        LEGEND,
        [COMPAT, LICENSE],
    )

    # Projects ------------------------------------------------------------------------
    title(projects, "Projects", "One row per project. Owner, status and priority have drop-down lists. Grey columns calculate automatically.")
    names = ["#", "Project", "Owner", "Status", "Priority", "Start", "Due", "Tasks", "Done", "Progress", "Days left", "Health", "Notes"]
    header(projects, 4, names, [5, 30, 14, 13, 10, 11, 11, 8, 8, 10, 10, 13, 38])
    body(projects, P_FIRST, P_LAST, len(names), calc_cols=(1, 8, 9, 10, 11, 12), date_cols=(6, 7), formats={10: "0%", 11: "0"})
    tr = f"Tasks!$B${T_FIRST}:$B${T_LAST}"
    ts = f"Tasks!$D${T_FIRST}:$D${T_LAST}"
    for r in range(P_FIRST, P_LAST + 1):
        projects[f"A{r}"] = f'=IF(B{r}="","",ROW()-{P_FIRST - 1})'
        projects[f"H{r}"] = f'=IF(B{r}="","",COUNTIF({tr},B{r}))'
        projects[f"I{r}"] = f'=IF(B{r}="","",COUNTIFS({tr},B{r},{ts},"Done"))'
        projects[f"J{r}"] = f'=IF(OR(B{r}="",H{r}=0),"",I{r}/H{r})'
        projects[f"K{r}"] = f'=IF(OR(B{r}="",G{r}="",D{r}="Done"),"",G{r}-TODAY())'
        projects[f"L{r}"] = f'=IF(B{r}="","",IF(D{r}="Done","Done",IF(G{r}="","No due date",IF(G{r}<TODAY(),"Overdue",IF(G{r}-TODAY()<=7,"Due soon","On track")))))'
    validation(projects, refs["People"], [f"C{P_FIRST}:C{P_LAST}"])
    validation(projects, refs["Status"], [f"D{P_FIRST}:D{P_LAST}"])
    validation(projects, refs["Priority"], [f"E{P_FIRST}:E{P_LAST}"])
    hl = f"L{P_FIRST}:L{P_LAST}"
    fill_rule(projects, hl, f'$L{P_FIRST}="Overdue"', fill="F6D5D0", font=Font(color="8A1C12", bold=True))
    fill_rule(projects, hl, f'$L{P_FIRST}="Due soon"', fill="FBE6BD", font=Font(color="6E4500", bold=True))
    fill_rule(projects, hl, f'$L{P_FIRST}="On track"', fill=GREEN_TINT, font=Font(color="24543B"))
    fill_rule(projects, f"A{P_FIRST}:M{P_LAST}", f'$D{P_FIRST}="Done"', font=Font(color="8A938D"))
    projects.freeze_panes = "C5"
    print_setup(projects, f"A1:M{P_LAST}", title_rows="4:4")

    # Tasks ---------------------------------------------------------------------------
    title(tasks, "Tasks", "One row per task. Pick the project from the list so the task counts towards it. Set the status to Done and add the date when it's finished.")
    names = ["Task", "Project", "Owner", "Status", "Priority", "Start", "Due", "Done on", "Flag", "Notes"]
    header(tasks, 4, names, [38, 26, 14, 13, 10, 11, 11, 11, 11, 36])
    body(tasks, T_FIRST, T_LAST, len(names), calc_cols=(9,), date_cols=(6, 7, 8))
    for r in range(T_FIRST, T_LAST + 1):
        tasks[f"I{r}"] = f'=IF(A{r}="","",IF(D{r}="Done","",IF(G{r}="","",IF(G{r}<TODAY(),"Overdue",IF(G{r}-TODAY()<=2,"Due soon","")))))'
    validation(tasks, f"=Projects!$B${P_FIRST}:$B${P_LAST}", [f"B{T_FIRST}:B{T_LAST}"])
    validation(tasks, refs["People"], [f"C{T_FIRST}:C{T_LAST}"])
    validation(tasks, refs["Status"], [f"D{T_FIRST}:D{T_LAST}"])
    validation(tasks, refs["Priority"], [f"E{T_FIRST}:E{T_LAST}"])
    fl = f"I{T_FIRST}:I{T_LAST}"
    fill_rule(tasks, fl, f'$I{T_FIRST}="Overdue"', fill="F6D5D0", font=Font(color="8A1C12", bold=True))
    fill_rule(tasks, fl, f'$I{T_FIRST}="Due soon"', fill="FBE6BD", font=Font(color="6E4500", bold=True))
    fill_rule(tasks, f"A{T_FIRST}:J{T_LAST}", f'$D{T_FIRST}="Done"', font=Font(color="8A938D", strike=True))
    tasks.freeze_panes = "B5"
    print_setup(tasks, f"A1:J{T_LAST}", title_rows="4:4")

    # Timeline ------------------------------------------------------------------------
    weeks = 26
    title(timeline, "Project timeline", "Bars come from each project's start and due dates on the Projects sheet. The current week is shaded yellow.")
    timeline["A3"] = "First week starts on (a Monday):"
    timeline["A3"].font = f(10, bold=True)
    timeline["D3"] = "=TODAY()-WEEKDAY(TODAY(),3)"
    timeline["D3"].fill = SETTING_FILL
    timeline["D3"].number_format = DATE
    timeline["D3"].border = BOX
    timeline["E3"] = "← shows this week by default; type a Monday's date to fix it"
    timeline["E3"].font = f(9, color="4A554F", italic=True)
    header(timeline, 5, ["Project", "Owner", "Start", "Due", "Status"], [28, 12, 10, 10, 12])
    first_week_col = 6
    for i in range(weeks):
        col = first_week_col + i
        letter = get_column_letter(col)
        prev = get_column_letter(col - 1)
        c = timeline.cell(row=5, column=col, value="=$D$3" if i == 0 else f"={prev}5+7")
        c.number_format = "d mmm"
        c.font = f(8.5, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        m = timeline.cell(row=4, column=col, value=f"={letter}5" if i == 0 else f'=IF(MONTH({letter}5)<>MONTH({prev}5),{letter}5,"")')
        m.number_format = "mmm yyyy"
        m.font = f(9, bold=True, color=GREEN)
        timeline.column_dimensions[letter].width = 6.5
    last_col = get_column_letter(first_week_col + weeks - 1)
    for i, r in enumerate(range(6, 6 + (P_LAST - P_FIRST + 1))):
        pr_ = P_FIRST + i
        timeline[f"A{r}"] = f'=IF(Projects!B{pr_}="","",Projects!B{pr_})'
        timeline[f"B{r}"] = f'=IF(Projects!C{pr_}="","",Projects!C{pr_})'
        timeline[f"C{r}"] = f'=IF(Projects!F{pr_}="","",Projects!F{pr_})'
        timeline[f"D{r}"] = f'=IF(Projects!G{pr_}="","",Projects!G{pr_})'
        timeline[f"E{r}"] = f'=IF(Projects!D{pr_}="","",Projects!D{pr_})'
        for col in range(1, first_week_col + weeks):
            c = timeline.cell(row=r, column=col)
            c.border = BOX
            c.font = f(9.5, color="4A554F")
            if col < first_week_col:
                c.fill = CALC_FILL
            if col in (3, 4):
                c.number_format = DATE
        timeline.row_dimensions[r].height = 18
    last_row = 6 + (P_LAST - P_FIRST)
    grid = f"F6:{last_col}{last_row}"
    bar = "AND(ISNUMBER($C6),ISNUMBER($D6),F$5<=$D6,F$5+6>=$C6)"
    fill_rule(timeline, grid, f'AND({bar},$E6="Done")', fill="B9C9BE")
    fill_rule(timeline, grid, bar, fill=GREEN)
    fill_rule(timeline, grid, "AND(F$5<=TODAY(),F$5+6>=TODAY())", fill="FFF2CC")
    fill_rule(timeline, f"F5:{last_col}5", "AND(F$5<=TODAY(),F$5+6>=TODAY())", fill="C9A227")
    timeline.freeze_panes = "F6"
    print_setup(timeline, f"A1:{last_col}{last_row}")

    # Monthly overview ----------------------------------------------------------------
    title(month, "Monthly overview", "Type any date in the yellow cell to show that month. Each day shows open tasks due and the first project due.")
    month["A3"] = "Month to show:"
    month["A3"].font = f(10, bold=True)
    month["B3"] = "=TODAY()"
    month["B3"].fill = SETTING_FILL
    month["B3"].number_format = "mmmm yyyy"
    month["B3"].font = f(11, bold=True)
    month["B3"].border = BOX
    month["C3"] = "← shows this month by default"
    month["C3"].font = f(9, color="4A554F", italic=True)
    first, cells = month_grid(month, "$B$3", 5)
    td = f"Tasks!$G${T_FIRST}:$G${T_LAST}"
    for r, col in cells:
        a = f"{get_column_letter(col)}{r}"
        count = f'COUNTIFS({td},{a},{ts},"<>Done")'
        month.cell(row=r + 1, column=col, value=f'=IF({count}=0,"",{count}&IF({count}=1," task due"," tasks due"))')
        month.cell(row=r + 2, column=col, value=f'=IFERROR("Due: "&INDEX(Projects!$B${P_FIRST}:$B${P_LAST},MATCH({a},Projects!$G${P_FIRST}:$G${P_LAST},0)),"")')
    month.column_dimensions["H"].width = 3
    month.column_dimensions["I"].width = 26
    month.column_dimensions["J"].width = 10
    month["I5"] = "This month"
    month["I5"].font = f(10, bold=True, color="FFFFFF")
    month["I5"].fill = HEAD_FILL
    month["J5"].fill = HEAD_FILL
    end = f"EOMONTH({first},0)"
    pd_ = f"Projects!$G${P_FIRST}:$G${P_LAST}"
    tdone = f"Tasks!$H${T_FIRST}:$H${T_LAST}"
    tname = f"Tasks!$A${T_FIRST}:$A${T_LAST}"
    stats = [
        ("Projects due", f'=COUNTIFS({pd_},">="&{first},{pd_},"<="&{end})'),
        ("Tasks due", f'=COUNTIFS({td},">="&{first},{td},"<="&{end})'),
        ("Tasks completed", f'=COUNTIFS({tdone},">="&{first},{tdone},"<="&{end})'),
        ("Tasks overdue today", f'=COUNTIFS({td},"<"&TODAY(),{ts},"<>Done",{tname},"<>")'),
    ]
    for i, (label, formula) in enumerate(stats):
        r = 6 + i
        month[f"I{r}"] = label
        month[f"I{r}"].font = f(10)
        month[f"J{r}"] = formula
        month[f"J{r}"].fill = CALC_FILL
        month[f"J{r}"].font = f(11, bold=True, color=GREEN)
        for c in (f"I{r}", f"J{r}"):
            month[c].border = BOX
    for label, top, rows in (("Focus for this month", 12, 5), ("Notes", 19, 10)):
        month[f"I{top}"] = label
        month[f"I{top}"].font = f(10, bold=True, color=GREEN)
        month.merge_cells(f"I{top + 1}:J{top + rows}")
        cell = month[f"I{top + 1}"]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        for rr in range(top + 1, top + rows + 1):
            for cc in ("I", "J"):
                month[f"{cc}{rr}"].border = BOX
    print_setup(month, "A1:J23")

    # Weekly checklist ----------------------------------------------------------------
    from printables import CHECKLIST

    title(weekly, "Weekly checklist", "Write the date each week starts in the yellow row, then choose ✓ for each step you complete.")
    header(weekly, 4, ["#", "Step"] + [f"Week {i}" for i in range(1, 9)], [4, 62] + [10] * 8)
    weekly["B5"] = "Week starting:"
    weekly["B5"].font = f(10, bold=True)
    for i in range(8):
        c = weekly.cell(row=5, column=3 + i)
        c.fill = SETTING_FILL
        c.number_format = DATE
        c.border = BOX
    r = 6
    n = 1
    item_rows = []
    for group, items in CHECKLIST:
        g = weekly.cell(row=r, column=2, value=group)
        g.font = f(10, bold=True, color=GREEN)
        g.fill = TINT_FILL
        for col in range(1, 11):
            weekly.cell(row=r, column=col).fill = TINT_FILL
        r += 1
        for item in items:
            weekly.cell(row=r, column=1, value=n).font = f(10, color="4A554F")
            s = weekly.cell(row=r, column=2, value=item)
            s.font = f(10)
            s.alignment = Alignment(wrap_text=True, vertical="center")
            for col in range(1, 11):
                weekly.cell(row=r, column=col).border = BOX
                if col >= 3:
                    weekly.cell(row=r, column=col).alignment = Alignment(horizontal="center", vertical="center")
                    weekly.cell(row=r, column=col).font = f(12, bold=True, color=GREEN)
            weekly.row_dimensions[r].height = 20
            item_rows.append(r)
            r += 1
            n += 1
    total = n - 1
    validation(weekly, '"✓"', [f"C{item_rows[0]}:J{item_rows[-1]}"])
    weekly.cell(row=r, column=2, value="Steps completed").font = f(10, bold=True)
    for i in range(8):
        col = get_column_letter(3 + i)
        c = weekly.cell(row=r, column=3 + i, value=f'=COUNTIF({col}{item_rows[0]}:{col}{item_rows[-1]},"✓")&" of {total}"')
        c.fill = CALC_FILL
        c.font = f(9.5, bold=True, color=GREEN)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    weekly.freeze_panes = "C6"
    print_setup(weekly, f"A1:J{r}")

    if sample:
        _sample_tracker(projects, tasks, weekly, item_rows)
        projects.print_area = f"A1:M{P_FIRST + 8}"
        tasks.print_area = f"A1:J{T_FIRST + 12}"
        timeline.print_area = f"A1:{last_col}14"

    for ws in (projects, tasks, timeline, month, weekly, lists):
        ws.sheet_view.zoomScale = 100
    wb.active = 0
    wb.save(path)
    return [ws.title for ws in wb.worksheets]


CHANNELS = ["Blog", "Newsletter", "Instagram", "LinkedIn", "YouTube", "TikTok", "Podcast", "Website", "Other"]
FORMATS = ["Article", "Short post", "Carousel", "Video", "Short video", "Email", "Episode", "Livestream", "Other"]
C_STATUS = ["Idea", "Drafting", "In review", "Scheduled", "Published", "Skipped"]
I_STATUS = ["New", "Planned", "Used", "Dropped"]
C_FIRST, C_LAST = 5, 404


def content_calendar(path: Path, sample: bool = False) -> list[str]:
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    start = wb.active
    start.title = "Start here"
    cal = wb.create_sheet("Calendar")
    month = wb.create_sheet("Month view")
    ideas = wb.create_sheet("Ideas")
    lists = wb.create_sheet("Lists")

    refs = lists_sheet(
        lists,
        [
            ("Channels", CHANNELS, 20, "Rename or add the places you publish."),
            ("Formats", FORMATS, 15, None),
            ("Status", C_STATUS, 8, "Keep “Published”, “Scheduled” and “Skipped”: the colors and flags use them."),
            ("People", ["Me"], 20, None),
            ("Idea status", I_STATUS, 6, None),
        ],
    )
    start_sheet(
        start,
        "Content Calendar",
        "Plan what you'll publish, where and when. Add each piece of content on the Calendar sheet, see the month at a glance on the Month view, "
        "and keep ideas that don't have a date yet on the Ideas sheet.",
        [
            "On the Lists sheet, edit the channels to match where you publish (blog, newsletter, social accounts and so on).",
            "On the Calendar sheet, add one row per piece of content with its publish date, channel, title and format.",
            "Move the status along as you work: Idea, Drafting, In review, Scheduled, Published. Scheduled rows turn blue and published rows turn green.",
            "Watch the Flag column: it warns about past dates that aren't published, drafts that are late, and anything due in the next 3 days.",
            "Open the Month view and type any date in the yellow cell to see that month's plan.",
        ],
        LEGEND,
        [COMPAT, LICENSE],
    )

    title(cal, "Content calendar", "One row per piece of content. Day and Week fill in from the date. Rows don't need to be in date order, but sorting by date helps.")
    names = ["Publish date", "Day", "Week", "Channel", "Title or topic", "Format", "Status", "Flag", "Owner", "Draft due", "Link or file", "Notes"]
    header(cal, 4, names, [13, 7, 7, 14, 40, 13, 12, 11, 13, 11, 30, 30])
    body(cal, C_FIRST, C_LAST, len(names), calc_cols=(2, 3, 8), date_cols=(1, 10), formats={2: "ddd", 3: "0"})
    for r in range(C_FIRST, C_LAST + 1):
        cal[f"B{r}"] = f'=IF(A{r}="","",A{r})'
        cal[f"C{r}"] = f'=IF(A{r}="","",WEEKNUM(A{r},21))'
        cal[f"H{r}"] = (
            f'=IF(A{r}="","",IF(OR(G{r}="Published",G{r}="Skipped"),"",IF(A{r}<TODAY(),"Past date",'
            f'IF(AND(ISNUMBER(J{r}),J{r}<TODAY(),G{r}<>"Scheduled"),"Draft late",IF(A{r}-TODAY()<=3,"Soon","")))))'
        )
    validation(cal, refs["Channels"], [f"D{C_FIRST}:D{C_LAST}"])
    validation(cal, refs["Formats"], [f"F{C_FIRST}:F{C_LAST}"])
    validation(cal, refs["Status"], [f"G{C_FIRST}:G{C_LAST}"])
    validation(cal, refs["People"], [f"I{C_FIRST}:I{C_LAST}"])
    rows = f"A{C_FIRST}:G{C_LAST}"
    fill_rule(cal, rows, f'$G{C_FIRST}="Published"', fill=GREEN_TINT)
    fill_rule(cal, rows, f'$G{C_FIRST}="Scheduled"', fill="E1E9F2")
    fl = f"H{C_FIRST}:H{C_LAST}"
    fill_rule(cal, fl, f'$H{C_FIRST}="Past date"', fill="F6D5D0", font=Font(color="8A1C12", bold=True))
    fill_rule(cal, fl, f'$H{C_FIRST}="Draft late"', fill="FBE6BD", font=Font(color="6E4500", bold=True))
    fill_rule(cal, fl, f'$H{C_FIRST}="Soon"', font=Font(color="24543B", bold=True))
    cal.freeze_panes = "B5"
    print_setup(cal, f"A1:L{C_LAST}", title_rows="4:4")

    title(month, "Month view", "Type any date in the yellow cell to show that month. Each day shows the first planned item and how many more there are.")
    month["A3"] = "Month to show:"
    month["A3"].font = f(10, bold=True)
    month["B3"] = "=TODAY()"
    month["B3"].fill = SETTING_FILL
    month["B3"].number_format = "mmmm yyyy"
    month["B3"].font = f(11, bold=True)
    month["B3"].border = BOX
    month["C3"] = "← shows this month by default"
    month["C3"].font = f(9, color="4A554F", italic=True)
    first, cells = month_grid(month, "$B$3", 5)
    cd = f"Calendar!$A${C_FIRST}:$A${C_LAST}"
    for r, col in cells:
        a = f"{get_column_letter(col)}{r}"
        m = f"MATCH({a},{cd},0)"
        month.cell(row=r + 1, column=col, value=f'=IFERROR(INDEX(Calendar!$D${C_FIRST}:$D${C_LAST},{m})&": "&INDEX(Calendar!$E${C_FIRST}:$E${C_LAST},{m}),"")')
        month.cell(row=r + 2, column=col, value=f'=IF(COUNTIF({cd},{a})>1,"+"&(COUNTIF({cd},{a})-1)&" more","")')
    month.column_dimensions["H"].width = 3
    month.column_dimensions["I"].width = 26
    month.column_dimensions["J"].width = 10
    month["I5"] = "This month"
    month["I5"].font = f(10, bold=True, color="FFFFFF")
    month["I5"].fill = HEAD_FILL
    month["J5"].fill = HEAD_FILL
    end = f"EOMONTH({first},0)"
    cs = f"Calendar!$G${C_FIRST}:$G${C_LAST}"
    in_month = f'{cd},">="&{first},{cd},"<="&{end}'
    stats = [
        ("Planned", f"=COUNTIFS({in_month})"),
        ("Published", f'=COUNTIFS({in_month},{cs},"Published")'),
        ("Still to publish", f'=COUNTIFS({in_month},{cs},"<>Published",{cs},"<>Skipped")'),
        ("Past date, not published", f'=COUNTIFS({cd},"<"&TODAY(),{cs},"<>Published",{cs},"<>Skipped")'),
    ]
    for i, (label, formula) in enumerate(stats):
        r = 6 + i
        month[f"I{r}"] = label
        month[f"J{r}"] = formula
        month[f"J{r}"].fill = CALC_FILL
        month[f"J{r}"].font = f(11, bold=True, color=GREEN)
        for c in (f"I{r}", f"J{r}"):
            month[c].border = BOX
    month["I12"] = "Themes for this month"
    month["I12"].font = f(10, bold=True, color=GREEN)
    month.merge_cells("I13:J22")
    month["I13"].alignment = Alignment(wrap_text=True, vertical="top")
    for rr in range(13, 23):
        for cc in ("I", "J"):
            month[f"{cc}{rr}"].border = BOX
    print_setup(month, "A1:J23")

    title(ideas, "Ideas", "Collect ideas before they have a date. When one is scheduled, copy it to the Calendar sheet and set its status to Used.")
    names = ["Idea", "Channel", "Format", "Added on", "Use by", "Status", "Notes"]
    header(ideas, 4, names, [46, 14, 13, 11, 11, 11, 40])
    body(ideas, 5, 204, len(names), date_cols=(4, 5))
    validation(ideas, refs["Channels"], ["B5:B204"])
    validation(ideas, refs["Formats"], ["C5:C204"])
    validation(ideas, refs["Idea status"], ["F5:F204"])
    fill_rule(ideas, "A5:G204", '$F5="Used"', font=Font(color="8A938D"))
    fill_rule(ideas, "A5:G204", '$F5="Dropped"', font=Font(color="8A938D", strike=True))
    ideas.freeze_panes = "B5"
    print_setup(ideas, "A1:G204", title_rows="4:4")

    if sample:
        _sample_calendar(cal, ideas)
        cal.print_area = f"A1:L{C_FIRST + 15}"
        ideas.print_area = "A1:G10"

    wb.active = 0
    wb.save(path)
    return [ws.title for ws in wb.worksheets]


# ---------------------------------------------------------------------------------------
# Sample entries, used only to render the website previews. Customers get blank files.


def _monday(today):
    return today - dt.timedelta(days=today.weekday())


def _sample_tracker(projects, tasks, weekly, item_rows):
    today = dt.date.today()
    mon = _monday(today)
    d = lambda n: mon + dt.timedelta(days=n)  # noqa: E731
    rows = [
        ("Website refresh", "Me", "In progress", "High", d(-21), d(24), "New home page and pricing copy"),
        ("Spring newsletter series", "Sam", "In progress", "Medium", d(-7), d(10), "Four issues"),
        ("Client onboarding pack", "Me", "Not started", "Medium", d(14), d(49), ""),
        ("Office move", "Alex", "Done", "Low", d(-42), d(-4), "Finished early"),
        ("Q3 planning", "Me", "Waiting", "High", d(-3), d(3), "Waiting on budget figures"),
    ]
    for i, row in enumerate(rows):
        r = P_FIRST + i
        for col, v in zip("BCDEFGM", row):
            projects[f"{col}{r}"] = v
    trows = [
        ("Write new home page copy", "Website refresh", "Me", "Done", "High", d(-21), d(-10), d(-11)),
        ("Choose photos", "Website refresh", "Me", "In progress", "Medium", d(-9), d(1), None),
        ("Review pricing page with Sam", "Website refresh", "Sam", "Not started", "High", d(2), d(5), None),
        ("Publish changes", "Website refresh", "Me", "Not started", "High", d(20), d(24), None),
        ("Draft issue 1", "Spring newsletter series", "Sam", "Done", "Medium", d(-7), d(-2), d(-3)),
        ("Draft issue 2", "Spring newsletter series", "Sam", "In progress", "Medium", d(-1), d(-1), None),
        ("Set up sending schedule", "Spring newsletter series", "Me", "Not started", "Low", d(3), d(7), None),
        ("Collect last year's figures", "Q3 planning", "Me", "Done", "High", d(-3), d(-1), d(-1)),
        ("Agree budget", "Q3 planning", "Alex", "Waiting", "High", d(0), d(2), None),
        ("Pack and label boxes", "Office move", "Alex", "Done", "Low", d(-42), d(-10), d(-12)),
    ]
    for i, row in enumerate(trows):
        r = T_FIRST + i
        for col, v in zip("ABCDEFGH", row):
            if v is not None:
                tasks[f"{col}{r}"] = v
    for i in range(3):
        weekly.cell(row=5, column=3 + i, value=mon - dt.timedelta(days=7 * (2 - i)))
    for k, r in enumerate(item_rows):
        for i in range(3):
            if (k + i) % 5 or i < 2:
                weekly.cell(row=r, column=3 + i, value="✓")


def _sample_calendar(cal, ideas):
    today = dt.date.today()
    first = today.replace(day=1)
    d = lambda n: first + dt.timedelta(days=n)  # noqa: E731
    rows = [
        (d(1), "Blog", "How we plan a month of posts", "Article", "Published", "Me", d(-2)),
        (d(3), "Newsletter", "Monthly update", "Email", "Published", "Sam", d(1)),
        (d(6), "Instagram", "Behind the scenes: workspace", "Carousel", "Published", "Me", None),
        (d(8), "LinkedIn", "Three lessons from our last launch", "Short post", "Published", "Me", None),
        (d(10), "YouTube", "Desk setup walk-through", "Video", "Scheduled", "Alex", d(6)),
        (d(13), "Blog", "Weekly review template, explained", "Article", "In review", "Me", d(9)),
        (d(15), "Newsletter", "Mid-month tips", "Email", "Drafting", "Sam", d(12)),
        (d(17), "Instagram", "Customer question of the week", "Short post", "Idea", "Me", None),
        (d(20), "Podcast", "Episode 4: planning on paper", "Episode", "Drafting", "Alex", d(16)),
        (d(22), "LinkedIn", "What a good brief includes", "Short post", "Idea", "Me", None),
        (d(24), "Blog", "Content calendar FAQ", "Article", "Idea", "Me", d(20)),
        (d(27), "Newsletter", "Month in review", "Email", "Idea", "Sam", d(24)),
        (d(13), "Instagram", "Template preview reel", "Short video", "Scheduled", "Alex", None),
    ]
    for i, row in enumerate(rows):
        r = C_FIRST + i
        for col, v in zip("ADEFGIJ", row):
            if v is not None:
                cal[f"{col}{r}"] = v
    for i, row in enumerate([
        ("Interview a customer about their weekly routine", "Blog", "Article", today, None, "New"),
        ("Short video: printing tips", "YouTube", "Short video", today, None, "Planned"),
        ("Checklist carousel", "Instagram", "Carousel", today, None, "New"),
    ]):
        for col, v in zip("ABCDEF", row):
            if v is not None:
                ideas[f"{col}{5 + i}"] = v
